import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)   # Slate 700
    PRIMARY = RGBColor(59, 130, 246)     # Blue 500
    CYAN = RGBColor(6, 182, 212)        # Cyan 500
    EMERALD = RGBColor(16, 185, 129)    # Emerald 500
    ROSE = RGBColor(244, 63, 94)        # Rose 500
    PURPLE = RGBColor(168, 85, 247)     # Purple 500
    AMBER = RGBColor(245, 158, 11)      # Amber 500
    TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    TEXT_DIM = RGBColor(100, 116, 139)   # Slate 500

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    def add_header(slide, tag, title, subtitle):
        set_slide_background(slide)
        
        # Tag pill
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN
        
        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.7))
        tf2 = t_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT

        # Subtitle
        s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.5))
        tf3 = s_box.text_frame
        tf3.word_wrap = True
        p3 = tf3.paragraphs[0]
        p3.text = subtitle
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Hero)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative top glow card
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    glow.fill.solid()
    glow.fill.fore_color.rgb = PRIMARY
    glow.line.fill.background()

    # Title box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "NEXT-GENERATION EDTECH PLATFORM"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.space_after = Pt(12)

    p1 = tf.add_paragraph()
    p1.text = "QUIZZY"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Gamified Active Recall & AI-Powered MCQ Exam Mastery"
    p2.font.size = Pt(22)
    p2.font.color.rgb = PRIMARY
    p2.space_after = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "Bridging the gap between passive content consumption and active knowledge retention through automated AI question extraction, spaced repetition 3D flashcards, and real-time competitive analytics."
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

    # Bottom pill badges
    badges = [
        ("AI Generation", "Gemini 3.5 Flash-Lite"),
        ("Live Deployment", "quizzy-su9d.onrender.com"),
        ("Core Tech", "React 18 + Express + PostgreSQL"),
        ("Architecture", "Full-Stack TypeScript Monorepo")
    ]
    card_w = Inches(2.65)
    gap = Inches(0.24)
    start_x = Inches(1.0)
    y_pos = Inches(5.6)
    for i, (b_label, b_val) in enumerate(badges):
        x = start_x + i * (card_w + gap)
        add_card(s1, x, y_pos, card_w, Inches(1.1))
        tb = s1.shapes.add_textbox(x + Inches(0.15), y_pos + Inches(0.15), card_w - Inches(0.3), Inches(0.8))
        ptf = tb.text_frame
        ptf.word_wrap = True
        pp1 = ptf.paragraphs[0]
        pp1.text = b_label.upper()
        pp1.font.size = Pt(10)
        pp1.font.bold = True
        pp1.font.color.rgb = CYAN
        
        pp2 = ptf.add_paragraph()
        pp2.text = b_val
        pp2.font.size = Pt(12)
        pp2.font.bold = True
        pp2.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 2: The Problem Statement
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Market Realities & Academic Pain Points", "The Problem: Why Traditional Studying Fails Students", 
               "Modern higher-education exams demand rigorous retrieval practice, yet students rely on obsolete study habits.")
    
    problems = [
        ("01 / Passive Cramming Illusion",
         "The 'Illusion of Competence'",
         "Re-reading lecture slides and highlighting textbooks gives students false confidence. Cognitive science proves passive review retains less than 15% of information after 48 hours without active testing.",
         ROSE),
        ("02 / The High-Friction Question Gap",
         "Lack of Relevant Practice MCQs",
         "Professors teach from diverse textbooks, but finding quality MCQs with reliable answer keys and explanatory rationale is arduous, scattered, or locked behind expensive paywalls.",
         AMBER),
        ("03 / The Forgotten Mistake Black Hole",
         "Mistakes Never Turned Into Mastery",
         "When students complete paper mock tests or ad-hoc quizzes, incorrect answers are checked once and forgotten. There is zero systematic loop to re-test weaknesses until exam day.",
         PURPLE),
        ("04 / The Study Motivation Cliff",
         "Solo Fatigue & Zero Engagement",
         "Traditional revision feels isolating and monotonous. Without visible streaks, level progression, or peer comparison, students succumb to procrastination and exam anxiety.",
         CYAN)
    ]

    col_w = Inches(2.75)
    col_gap = Inches(0.20)
    col_x = Inches(0.8)
    col_y = Inches(2.2)
    card_h = Inches(4.7)

    for i, (num, headline, body, accent) in enumerate(problems):
        x = col_x + i * (col_w + col_gap)
        add_card(s2, x, col_y, col_w, card_h)
        
        # Accent top bar inside card
        bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, col_y, col_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(x + Inches(0.2), col_y + Inches(0.3), col_w - Inches(0.4), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_num = tf.paragraphs[0]
        p_num.text = num
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = accent
        p_num.space_after = Pt(10)

        p_h = tf.add_paragraph()
        p_h.text = headline
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_LIGHT
        p_h.space_after = Pt(12)

        p_b = tf.add_paragraph()
        p_b.text = body
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 3: The Solution - The Quizzy Ecosystem
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Transformative Learning Model", "The Solution: Quizzy's 3-Pillar Learning Ecosystem", 
               "An integrated platform blending active retrieval, cognitive spaced repetition, and dopamine-driven gamification.")

    pillars = [
        ("PILLAR 1", "Active Testing Engine", 
         "Instant Retrieval Practice",
         ["Timed exam mode with negative marking simulation", 
          "Curated semester-wise subjects (Calculus, Physics, OS, etc.)", 
          "Granular instant answer verification with full reasoning", 
          "Zero-friction guest or authenticated student profiles"],
         PRIMARY),
        ("PILLAR 2", "3D Revision Hub", 
         "Cognitive Spaced Repetition",
         ["Every incorrect quiz answer automatically auto-saved", 
          "Interactive 3D flip cards with genuine rationale & briefing", 
          "Mastery progression tracking (Learning -> Mastered)", 
          "Custom mistake search and subject filter categorization"],
         EMERALD),
        ("PILLAR 3", "AI Question Factory", 
         "Instant Curriculum Digestion",
         ["Powered by Gemini 3.5 Flash-Lite API", 
          "Converts raw syllabus PDFs, lecture notes, or text in seconds", 
          "Strict JSON schema enforcement with automated rationale generation", 
          "Admin review & 1-click batch publishing to live database"],
         PURPLE)
    ]

    p_w = Inches(3.75)
    p_gap = Inches(0.24)
    p_x = Inches(0.8)
    p_y = Inches(2.2)
    p_h = Inches(4.7)

    for i, (tag, title, sub, bullets, accent) in enumerate(pillars):
        x = p_x + i * (p_w + p_gap)
        add_card(s3, x, p_y, p_w, p_h)

        # Header tag
        bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, p_y, p_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(x + Inches(0.25), p_y + Inches(0.3), p_w - Inches(0.5), p_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = tag
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = accent
        p_t.space_after = Pt(6)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(19)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = Pt(14)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_LIGHT
            pb.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: Interactive MCQ Testing Engine
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Feature Deep Dive 01", "The MCQ Engine: Exam Simulation at Professional Scale", 
               "Engineered to replicate actual university and competitive exam pressures while providing micro-learning feedback.")

    # 2-column layout: left card (Core Capabilities), right card (User Experience Highlights)
    w_2col = Inches(5.7)
    left_x1 = Inches(0.8)
    left_x2 = Inches(6.8)
    c_y = Inches(2.2)
    c_h = Inches(4.7)

    add_card(s4, left_x1, c_y, w_2col, c_h)
    tb1 = s4.shapes.add_textbox(left_x1 + Inches(0.3), c_y + Inches(0.3), w_2col - Inches(0.6), c_h - Inches(0.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "TESTING SPECIFICATIONS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(8)

    items1 = [
        ("Configurable Time Constraints", "Real-time countdown timer with automated submission safeguards to build authentic test endurance."),
        ("Multi-Option Rationale Mapping", "Every option (A, B, C, D) is linked to rigorous mathematical or theoretical briefing citations."),
        ("Negative Marking & Accuracy Logic", "Prepares students for real standardized competitive exams by penalizing blind guesswork."),
        ("Dynamic Question Randomization", "Questions and choices shuffle dynamically to prevent location-based rote memorization."),
        ("Session State Persistence", "Answers and ongoing attempts survive browser reloads and network reconnects.")
    ]
    for h, b in items1:
        ph = tf1.add_paragraph()
        ph.text = "▶  " + h
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_LIGHT
        pb = tf1.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(8)

    add_card(s4, left_x2, c_y, w_2col, c_h)
    tb2 = s4.shapes.add_textbox(left_x2 + Inches(0.3), c_y + Inches(0.3), w_2col - Inches(0.6), c_h - Inches(0.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "STUDENT EXPERIENCE & FEEDBACK"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.space_after = Pt(8)

    items2 = [
        ("Instant Post-Test Breakdown", "Granular performance report generated within milliseconds of quiz completion with score percentage."),
        ("Color-Coded Option Analysis", "Clear visual distinction between selected wrong answer (red) and actual correct choice (green)."),
        ("One-Click Question Bookmarking", "Flag complex problems during the exam to review during post-test debriefing."),
        ("Seamless Next-Quiz Recommendation", "AI suggests the next topic based on weakest subject accuracy from the session."),
        ("Semester-Aligned Filtering", "Students filter content specifically for their enrolled university semester (Sem 1 to 8).")
    ]
    for h, b in items2:
        ph = tf2.add_paragraph()
        ph.text = "★  " + h
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_LIGHT
        pb = tf2.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 5: 3D Revision Hub & Spaced Repetition
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Feature Deep Dive 02", "3D Revision Hub: Automated Mistakes to Mastery", 
               "Closing the learning loop with interactive flashcards that convert wrong exam answers into long-term knowledge.")

    # 3 horizontal cards
    f_w = Inches(3.75)
    f_gap = Inches(0.24)
    f_x = Inches(0.8)
    f_y = Inches(2.2)
    f_h = Inches(4.7)

    rev_features = [
        ("AUTOMATIC INGESTION", "Zero Manual Logging", 
         "Every time a student picks an incorrect option in any quiz, Quizzy's background event listener immediately captures the mistake into PostgreSQL with question metadata, selected answer, and explanation.",
         [
             "No tedious manual copy-pasting",
             "Dual-synced: LocalStorage + PostgreSQL",
             "Preserves historical timestamp & attempts"
         ],
         ROSE),
        ("INTERACTIVE 3D CARDS", "Tactile Micro-Learning",
         "Features smooth 3D CSS perspective card flips. The front presents the problem statement and the student's mistaken choice; the back reveals the correct answer and genuine briefing explanation.",
         [
             "3D flip interaction engages active recall",
             "Direct citation from question database",
             "Keyboard navigation (Flip, Next, Previous)"
         ],
         CYAN),
        ("MASTERY LIFECYCLE", "Leitner Spaced Retention",
         "Students mark cards as 'Mastered' as they re-test themselves. Mastered cards move out of active rotation but remain accessible in revision archives, reinforcing long-term memory.",
         [
             "Visual progress bar (% Mastered)",
             "Filter by 'Needs Review' vs 'Mastered'",
             "Prevents recency bias in study sessions"
         ],
         EMERALD)
    ]

    for i, (tag, title, desc, points, accent) in enumerate(rev_features):
        x = f_x + i * (f_w + f_gap)
        add_card(s5, x, f_y, f_w, f_h)

        bar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, f_y, f_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s5.shapes.add_textbox(x + Inches(0.25), f_y + Inches(0.3), f_w - Inches(0.5), f_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(6)

        pt = tf.add_paragraph()
        pt.text = title
        pt.font.size = Pt(18)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        pt.space_after = Pt(8)

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(12)

        for pt_item in points:
            pp = tf.add_paragraph()
            pp.text = "✔ " + pt_item
            pp.font.size = Pt(11)
            pp.font.bold = True
            pp.font.color.rgb = TEXT_LIGHT
            pp.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: AI Question Factory (Gemini 3.5 Flash-Lite)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Feature Deep Dive 03", "AI Question Generator: Turning Courseware into Quizzes", 
               "Leveraging Google Gemini 3.5 Flash-Lite for high-speed, schema-validated curriculum question extraction.")

    # 4 horizontal step cards
    st_w = Inches(2.75)
    st_gap = Inches(0.20)
    st_x = Inches(0.8)
    st_y = Inches(2.2)
    st_h = Inches(4.7)

    steps = [
        ("STEP 01", "Upload / Prompt", 
         "Educators or admins paste lecture text, syllabus objectives, or upload PDF course materials into the Quizzy Admin Hub.",
         "Raw text, OCR snippets, or academic lecture outlines accepted."),
        ("STEP 02", "LLM Extraction", 
         "Google Gemini 3.5 Flash-Lite processes context in < 3 seconds with low latency and optimal token consumption economics.",
         "Generates 4 plausible options, distractors, and correct answer."),
        ("STEP 03", "Rationale & Schema", 
         "Server enforces strict JSON schema typing: Question text, 4 choices, correct answer index (0-3), and comprehensive briefing.",
         "Zero hallucinated format errors; invalid structures auto-rejected."),
        ("STEP 04", "Publish & Database", 
         "Admin inspects, tweaks wording or options if desired, and publishes directly into PostgreSQL subject tables with 1 click.",
         "Available immediately to thousands of students across the globe.")
    ]

    for i, (step_tag, step_title, step_desc, step_sub) in enumerate(steps):
        x = st_x + i * (st_w + st_gap)
        add_card(s6, x, st_y, st_w, st_h)

        bar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, st_y, st_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PRIMARY
        bar.line.fill.background()

        tb = s6.shapes.add_textbox(x + Inches(0.2), st_y + Inches(0.3), st_w - Inches(0.4), st_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = step_tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN
        p.space_after = Pt(6)

        pt = tf.add_paragraph()
        pt.text = step_title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        pt.space_after = Pt(10)

        pd = tf.add_paragraph()
        pd.text = step_desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(10)

        ps = tf.add_paragraph()
        ps.text = "💡 " + step_sub
        ps.font.size = Pt(10)
        ps.font.italic = True
        ps.font.color.rgb = TEXT_DIM

    # -------------------------------------------------------------
    # SLIDE 7: Gamification & Dopamine-Driven Retention
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Feature Deep Dive 04", "Gamification & Community: Making Studying Addictive", 
               "Integrating behavioral psychology and habit loops to turn daily study sessions into high-engagement rituals.")

    w_2col = Inches(5.7)
    add_card(s7, left_x1, c_y, w_2col, c_h)
    tb1 = s7.shapes.add_textbox(left_x1 + Inches(0.3), c_y + Inches(0.3), w_2col - Inches(0.6), c_h - Inches(0.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "XP PROGRESSION & TIER SYSTEM"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_after = Pt(8)

    tiers = [
        ("Beginner (0 - 200 XP)", "🌱 Starting out on fundamental quizzes"),
        ("Scholar (201 - 500 XP)", "📚 Developing consistent subject familiarity"),
        ("Expert (501 - 1,000 XP)", "🎓 High-accuracy multi-subject mastery"),
        ("Master (1,001 - 2,000 XP)", "⚔️ Competitive speed and near-zero error rates"),
        ("Legend (2,001+ XP)", "👑 Top 1% leaderboard dominance and elite streaks")
    ]
    for name, desc in tiers:
        pn = tf1.add_paragraph()
        pn.text = "• " + name
        pn.font.size = Pt(12)
        pn.font.bold = True
        pn.font.color.rgb = TEXT_LIGHT
        pd = tf1.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    add_card(s7, left_x2, c_y, w_2col, c_h)
    tb2 = s7.shapes.add_textbox(left_x2 + Inches(0.3), c_y + Inches(0.3), w_2col - Inches(0.6), c_h - Inches(0.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "BADGES, STREAKS & AVATARS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PURPLE
    p.space_after = Pt(8)

    gamify_items = [
        ("Daily Habit Streaks 🔥", "Dynamic streak tracking rewards consecutive days of study, minimizing exam procrastination."),
        ("Dynamic Achievement Badges 🏆", "Unlockable accolades: 'Speed Demon' (record completion), 'Sharp Shooter' (90%+ accuracy), 'Comeback Kid' (30%+ score jump)."),
        ("Custom Animated Avatars 👾", "Visual personality customization tied to user profile with persistent SVG and CSS avatar states."),
        ("Global & Subject Leaderboards 🌟", "Live competitive rankings inspire friendly rivalries across university batches and semesters.")
    ]
    for h, b in gamify_items:
        ph = tf2.add_paragraph()
        ph.text = "★ " + h
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_LIGHT
        pb = tf2.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_MUTED
        pb.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 8: Analytics & Weak-Spot Diagnostics
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Feature Deep Dive 05", "Analytics & Insights: Data-Driven Performance Clarity", 
               "Transforming raw attempt data into actionable intelligence so students study smarter, not just longer.")

    an_w = Inches(3.75)
    an_gap = Inches(0.24)
    an_x = Inches(0.8)
    an_y = Inches(2.2)
    an_h = Inches(4.7)

    an_cards = [
        ("DIAGNOSTIC MATRIX", "Subject Mastery Split",
         "Automatically categorizes subjects into Strong (≥70% accuracy) and Weak (<70% accuracy) so revision time is never wasted on topics already conquered.",
         PRIMARY),
        ("STUDY DURATION TELEMETRY", "Weekly Engagement Curve",
         "Aggregates total quiz minutes and active exam focus time into interactive charts with baseline domains to track real revision momentum day-by-day.",
         CYAN),
        ("ACCURACY RADAR", "Overall Competence Score",
         "Computes weighted accuracy percentages across all attempted quizzes, displaying lifetime questions solved, average score, and progression delta.",
         EMERALD)
    ]

    for i, (tag, title, desc, accent) in enumerate(an_cards):
        x = an_x + i * (an_w + an_gap)
        add_card(s8, x, an_y, an_w, an_h)

        bar = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, an_y, an_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s8.shapes.add_textbox(x + Inches(0.25), an_y + Inches(0.3), an_w - Inches(0.5), an_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(6)

        pt = tf.add_paragraph()
        pt.text = title
        pt.font.size = Pt(19)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        pt.space_after = Pt(12)

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 9: Technology Stack & Full-Stack Architecture
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Technical Engineering", "Full-Stack Architecture: Built for Scale & Speed", 
               "Modern TypeScript monorepo with production-grade separation of concerns, strict type-safety, and cloud elasticity.")

    stk_w = Inches(2.75)
    stk_gap = Inches(0.20)
    stk_x = Inches(0.8)
    stk_y = Inches(2.2)
    stk_h = Inches(4.7)

    stack_layers = [
        ("FRONTEND LAYER", "React 18 & Vite", 
         ["Vite 7.3 build pipeline",
          "Tailwind CSS & Shadcn/UI",
          "Wouter lightweight routing",
          "Framer Motion 3D animations",
          "Recharts visual analytics",
          "Lucide-React icon system"],
         PRIMARY),
        ("BACKEND API", "Node.js & Express", 
         ["TypeScript ES modules",
          "Zod schema validation",
          "Pino high-speed logger",
          "OpenAPI / Swagger specs",
          "Session & Passport Auth",
          "CORS & Security middlewares"],
         CYAN),
        ("DATA & PERSISTENCE", "PostgreSQL & Drizzle", 
         ["Neon Serverless PostgreSQL",
          "Drizzle ORM type-safe queries",
          "Foreign-key relational integrity",
          "Optimized composite indexing",
          "LocalStorage offline fallback",
          "Automated migration scripts"],
         EMERALD),
        ("CLOUD & AI OPS", "Render & Gemini", 
         ["Render Web Service hosting",
          "Continuous GitHub CD pipeline",
          "Google Gemini 3.5 Flash-Lite",
          "Automatic SSL & custom domains",
          "Zero-downtime rolling deploys",
          "Healthcheck & metric probes"],
         PURPLE)
    ]

    for i, (tag, title, bullets, accent) in enumerate(stack_layers):
        x = stk_x + i * (stk_w + stk_gap)
        add_card(s9, x, stk_y, stk_w, stk_h)

        bar = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, stk_y, stk_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s9.shapes.add_textbox(x + Inches(0.2), stk_y + Inches(0.3), stk_w - Inches(0.4), stk_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(6)

        pt = tf.add_paragraph()
        pt.text = title
        pt.font.size = Pt(17)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        pt.space_after = Pt(12)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(11)
            pb.font.color.rgb = TEXT_LIGHT
            pb.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 10: Competitive Advantage Matrix
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Market Differentiation", "Competitive Landscape: How Quizzy Outperforms Existing Tools", 
               "Unifying the active testing of Quizlet, the gamified excitement of Kahoot, and the scientific retention of Anki.")

    # Table layout
    rows = 6
    cols = 5
    table_shape = s10.shapes.add_table(rows, cols, Inches(0.8), Inches(2.3), Inches(11.733), Inches(4.5))
    table = table_shape.table

    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(2.1)
    table.columns[2].width = Inches(2.1)
    table.columns[3].width = Inches(2.1)
    table.columns[4].width = Inches(2.233)

    table_data = [
        ["Key Capabilities", "Quizzy (Our App)", "Quizlet", "Kahoot", "Anki"],
        ["Automated Mistake Tracking", "✅ Built-in & Instant", "❌ Manual Only", "❌ No Persistence", "❌ Manual Only"],
        ["AI Question Extraction", "✅ Gemini 3.5 Flash-Lite", "⚠️ Paid Tier Only", "❌ None / Basic", "❌ Third-party Plugin"],
        ["Authentic Exam Simulation", "✅ Timed + Negative Marking", "❌ Casual Multiple Choice", "❌ Fast-twitch trivia", "❌ Pure Flashcards"],
        ["Gamification & Streaks", "✅ XP, Badges & Avatars", "⚠️ Minimal Badges", "✅ High in-class only", "❌ Very Dry / Zero"],
        ["Zero-Paywall Core Learning", "✅ 100% Free & Open", "❌ Heavy Paywalls ($35+/yr)", "❌ Freemium / Schools", "✅ Open-source (Steep curve)"]
    ]

    for r_idx, row in enumerate(table_data):
        for c_idx, cell_value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_value
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = CYAN
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG
            else:
                if c_idx == 1:
                    p.font.bold = True
                    p.font.color.rgb = EMERALD
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(20, 35, 55)
                else:
                    p.font.color.rgb = TEXT_LIGHT
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = CARD_BG

    # -------------------------------------------------------------
    # SLIDE 11: Roadmap & Future Expansion
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Growth Strategy", "Product Roadmap: The Evolution into a Global Learning Hub", 
               "Strategic milestones expanding Quizzy from an exam companion into an AI-native education ecosystem.")

    rm_cards = [
        ("PHASE 1 (COMPLETED)", "Core Foundation", 
         ["High-accuracy MCQ engine", "PostgreSQL mistake syncing", "3D Flashcard Revision Hub", "Render production deployment"],
         EMERALD),
        ("PHASE 2 (NEAR-TERM)", "Social & Collaborative", 
         ["Multiplayer real-time quiz lobbies", "Classroom batch invitations", "Custom user-generated deck sharing", "Offline PWA mobile app support"],
         PRIMARY),
        ("PHASE 3 (LONG-TERM)", "Socratic AI Tutor", 
         ["Voice-guided conversational tutor", "Dynamic adaptive difficulty adjustment", "Predictive exam readiness index", "LMS (Canvas/Moodle) integration"],
         PURPLE)
    ]

    rm_w = Inches(3.75)
    rm_gap = Inches(0.24)
    rm_x = Inches(0.8)
    rm_y = Inches(2.2)
    rm_h = Inches(4.7)

    for i, (tag, title, bullets, accent) in enumerate(rm_cards):
        x = rm_x + i * (rm_w + rm_gap)
        add_card(s11, x, rm_y, rm_w, rm_h)

        bar = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, rm_y, rm_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = s11.shapes.add_textbox(x + Inches(0.25), rm_y + Inches(0.3), rm_w - Inches(0.5), rm_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(6)

        pt = tf.add_paragraph()
        pt.text = title
        pt.font.size = Pt(19)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT
        pt.space_after = Pt(14)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "➔ " + b
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_LIGHT
            pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 12: Conclusion & Live Demonstration
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)

    # Center card
    card_center = add_card(s12, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    
    tb = s12.shapes.add_textbox(Inches(2.0), Inches(1.6), Inches(9.333), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "REVOLUTIONIZING ACTIVE STUDYING"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.alignment = PP_ALIGN.CENTER
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Empowering Every Student to Master Any Subject"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(16)

    p2 = tf.add_paragraph()
    p2.text = "Quizzy transforms passive, stressful exam cramming into an engaging, measurable, and scientifically validated journey towards true academic mastery."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "🌐 Live Application: https://quizzy-su9d.onrender.com\n💻 Repository: github.com/JawadAhmedRiaz/Quizzy\n⚡ Powered by: Google Gemini 3.5 Flash-Lite & Neon PostgreSQL"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = PRIMARY
    p3.alignment = PP_ALIGN.CENTER

    out_path = "D:\\Quizzy-main\\Quizzy_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation successfully created at {out_path}")

if __name__ == "__main__":
    build_deck()
